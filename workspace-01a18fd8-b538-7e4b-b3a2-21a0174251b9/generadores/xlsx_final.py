# -*- coding: utf-8 -*-
"""Excel de normalizacion BD NIKE Inc. - Grupo 05.

Cuatro hojas, una por forma normal: 0FN, 1FN, 2FN y 3FN.
Formato:
  * banner morado con el nombre de la forma normal
  * banda verde "TABLA : NOMBRE" antes de cada tabla
  * cabecera azul oscuro con el nombre de la columna y su marca (PK) / (FK) / (PK/FK)
  * celdas de PK y FK en crema; datos normales en blanco
  * en 0FN y 1FN los telefonos y las celdas con varios valores van marcados
Sin notas ni comentarios: eso vive en las guias.

Contenido de cada hoja
  0FN  tabla general: 3 pedidos con todo mezclado, incluidos los productos y las
       cantidades suministradas y el estado del envio.
  1FN  la misma informacion con un solo valor por celda (una fila por linea).
  2FN  8 tablas: PEDIDO, DETALLE_VENTA, PRODUCTO (con su categoria), SUMINISTRO,
       DETALLE_SUMINISTRO, ALMACEN, LOCAL y STOCK_ALMACEN.
  3FN  las 10 tablas del modelo, con CLIENTE y CATEGORIA separadas.
"""
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import modelo_final as M

MORADO = "7030A0"
VERDE = "00B050"
AZUL = "1F4E79"
CREMA = "FFF2CC"
SALMON = "FDE9D9"
GRIS = "F2F2F2"
BORDE_CLARO = "BFBFBF"

CAJA = Border(*[Side(style="thin", color=BORDE_CLARO)] * 4)

F_BANNER = Font(name="Calibri", size=18, bold=True, italic=True, color="FFFFFF")
F_BANDA = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
F_CAB = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
F_DATO = Font(name="Calibri", size=10, color="1A1A1A")
F_DATO_B = Font(name="Calibri", size=10, bold=True, color="1A1A1A")
F_ROJO = Font(name="Calibri", size=10, bold=True, color="C00000")
F_NULL = Font(name="Calibri", size=10, italic=True, color="808080")


def banner(ws, fila, texto, ncols):
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ncols)
    c = ws.cell(row=fila, column=1, value=texto)
    c.font = F_BANNER
    c.fill = PatternFill("solid", fgColor=MORADO)
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[fila].height = 30
    return fila + 1


def banda(ws, fila, nombre, ncols):
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ncols)
    c = ws.cell(row=fila, column=1, value="TABLA : %s" % nombre)
    c.font = F_BANDA
    c.fill = PatternFill("solid", fgColor=VERDE)
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[fila].height = 22
    return fila + 1


def tabla(ws, fila, cabecera, filas, claves, anchos=None, rojo_cols=(), salmon_cols=()):
    """claves: True por columna cuando la celda va en crema (PK o FK)."""
    for j, texto in enumerate(cabecera, 1):
        c = ws.cell(row=fila, column=j, value=texto)
        c.font = F_CAB
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = CAJA
    ws.row_dimensions[fila].height = 30
    for i, datos in enumerate(filas, 1):
        for j, valor in enumerate(datos, 1):
            c = ws.cell(row=fila + i, column=j, value=valor)
            c.border = CAJA
            c.alignment = Alignment(horizontal="center", vertical="center")
            if valor == "NULL":
                c.font = F_NULL
                c.fill = PatternFill("solid", fgColor=GRIS)
            elif (j - 1) in rojo_cols:
                c.font = F_ROJO
                c.fill = PatternFill("solid", fgColor=SALMON)
            elif (j - 1) in salmon_cols:
                c.font = F_DATO
                c.fill = PatternFill("solid", fgColor=SALMON)
            elif claves[j - 1]:
                c.font = F_DATO_B
                c.fill = PatternFill("solid", fgColor=CREMA)
            else:
                c.font = F_DATO
        ws.row_dimensions[fila + i].height = 18
    if anchos:
        for j, w in enumerate(anchos, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
    return fila + len(filas) + 2


def cabecera_3fn(tabla):
    """Nombre de la columna con su marca: id_cliente (PK) / CLIENTE_id_cliente (FK)."""
    out = []
    for nombre, tipo, largo, marca, opc in M.TABLES[tabla]["cols"]:
        et = nombre
        if marca == "PK":
            et += " (PK)"
        elif marca == "PFK":
            et += " (PK/FK)"
        elif marca == "FK":
            et += " (FK)"
        out.append(et)
    return out


def columnas_clave(tabla):
    return [c[3] in ("PK", "FK", "PFK") for c in M.TABLES[tabla]["cols"]]


# ------------------------------------------------------------------ hoja 0FN
COLS_0FN = ["id_pedido", "fecha", "id_cliente", "nombre_cliente", "correo_electronico",
            "telefonos_cliente", "direccion_cliente", "metodo_pago", "igv", "estado_envio",
            "productos_comprados", "categorias", "tallas", "colores", "cantidades_vendidas",
            "precios_aplicados", "productos_suministrados", "cantidades_suministradas",
            "almacen", "stock (talla/color/actual/minimo)", "local", "direccion_local",
            "fecha_suministro", "estado_despacho"]
FILAS_0FN = [
    ["9001", "2026-10-05", "1", "Carlos Mendoza", "cmendoza@gmail.com", "987654321, 945112233",
     "Av. Javier Prado 1240, San Isidro - Lima", "Tarjeta de Credito", "188.96", "ENTREGADO",
     "1001 - Nike Air Zoom Pegasus 40, 1002 - Nike Air Force 1 07",
     "10 - Calzado Running, 20 - Calzado Urbano", "42, 41", "Negro, Blanco", "1, 1",
     "549.90, 499.90", "1001 - Nike Air Zoom Pegasus 40, 1002 - Nike Air Force 1 07",
     "20, 15", "CD Lima Norte", "42/Negro/85/15, 41/Blanco/120/20", "Nike Store Jockey Plaza",
     "Av. Javier Prado Este 4200, Surco - Lima", "2026-10-04", "ENTREGADO"],
    ["9002", "2026-10-06", "2", "Lucia Paredes", "lparedes@hotmail.com", "912345678",
     "Av. Ejercito 710, Cayma - Arequipa", "Yape", "89.98", "ENVIADO",
     "1002 - Nike Air Force 1 07", "20 - Calzado Urbano", "41", "Blanco", "1", "499.90",
     "1002 - Nike Air Force 1 07", "25", "CD Lima Norte", "41/Blanco/120/20",
     "Nike Store Mall Arequipa", "Av. Porongoche 500, Paucarpata - Arequipa", "2026-10-05",
     "EN RUTA"],
    ["9003", "2026-10-06", "1", "Carlos Mendoza", "cmendoza@gmail.com", "987654321, 945112233",
     "Av. Javier Prado 1240, San Isidro - Lima", "Tarjeta de Debito", "46.76", "PENDIENTE",
     "1003 - Camiseta Nike Dri-FIT Academy", "30 - Indumentaria Deportiva", "M", "Azul Marino",
     "2", "129.90", "1003 - Camiseta Nike Dri-FIT Academy", "40", "CD Trujillo Norte",
     "M/Azul Marino/200/30", "Nike Store Mallplaza Trujillo",
     "Av. America Oeste 750 - Trujillo", "2026-10-06", "PROGRAMADO"],
]


def hoja_0fn(wb):
    ws = wb.create_sheet("0FN")
    n = len(COLS_0FN)
    f = banner(ws, 1, "TABLA GENERAL / UNIVERSAL (SIN NORMALIZAR - 0FN)", n)
    f = banda(ws, f + 1, "TABLA GENERAL", n)
    claves = [i in (2, 5, 10, 11, 12, 13, 14, 15, 16, 17, 19) for i in range(n)]
    f = tabla(ws, f, COLS_0FN, FILAS_0FN, claves,
              anchos=[12, 12, 10, 17, 24, 22, 34, 17, 9, 13, 38, 40, 11, 16, 12, 17, 38, 15, 16,
                      26, 24, 34, 15, 14],
              rojo_cols=(5,), salmon_cols=(10, 11, 12, 13, 14, 15, 16, 17, 19))
    ws.freeze_panes = "B4"
    return ws


# ------------------------------------------------------------------ hoja 1FN
COLS_1FN = ["id_pedido", "fecha", "id_cliente", "nombre_cliente", "correo_electronico", "telefono",
            "telefono_secundario", "direccion", "metodo_pago", "estado_envio", "id_producto",
            "nombre_producto",
            "talla_elegida", "color_elegido", "cantidad_vendida", "precio_unitario_aplicado",
            "id_categoria", "nombre_categoria", "descripcion_categoria", "id_suministro",
            "fecha_suministro", "estado_despacho", "id_producto_suministrado",
            "producto_suministrado", "talla_suministro", "color_suministro",
            "cantidad_suministrada", "id_almacen", "nombre_almacen", "ubicacion", "stock_actual",
            "stock_minimo", "id_local", "nombre_local", "direccion_local"]
FILAS_1FN = [
    ["9001", "2026-10-05", "1", "Carlos Mendoza", "cmendoza@gmail.com", "987654321", "945112233",
     "Av. Javier Prado 1240, San Isidro - Lima", "Tarjeta de Credito", "ENTREGADO", "1001",
     "Nike Air Zoom Pegasus 40", "42", "Negro", 1, 549.90, "10", "Calzado Running",
     "Zapatillas de entrenamiento y maraton", "1", "2026-10-04", "ENTREGADO", "1001",
     "Nike Air Zoom Pegasus 40", "42", "Negro", 20, "100", "CD Lima Norte",
     "Av. Ramiro Priale Km 8 - Lima", 85, 15, "501", "Nike Store Jockey Plaza",
     "Av. Javier Prado Este 4200, Surco - Lima"],
    ["9001", "2026-10-05", "1", "Carlos Mendoza", "cmendoza@gmail.com", "987654321", "945112233",
     "Av. Javier Prado 1240, San Isidro - Lima", "Tarjeta de Credito", "ENTREGADO", "1002",
     "Nike Air Force 1 07", "41", "Blanco", 1, 499.90, "20", "Calzado Urbano",
     "Zapatillas de uso diario Sportswear", "1", "2026-10-04", "ENTREGADO", "1002",
     "Nike Air Force 1 07", "41", "Blanco", 15, "100", "CD Lima Norte",
     "Av. Ramiro Priale Km 8 - Lima", 120, 20, "501", "Nike Store Jockey Plaza",
     "Av. Javier Prado Este 4200, Surco - Lima"],
    ["9002", "2026-10-06", "2", "Lucia Paredes", "lparedes@hotmail.com", "912345678", "NULL",
     "Av. Ejercito 710, Cayma - Arequipa", "Yape", "ENVIADO", "1002", "Nike Air Force 1 07", "41", "Blanco",
     1, 499.90, "20", "Calzado Urbano", "Zapatillas de uso diario Sportswear", "2", "2026-10-05",
     "EN RUTA", "1002", "Nike Air Force 1 07", "41", "Blanco", 25, "100", "CD Lima Norte",
     "Av. Ramiro Priale Km 8 - Lima", 120, 20, "502", "Nike Store Mall Arequipa",
     "Av. Porongoche 500, Paucarpata - Arequipa"],
    ["9003", "2026-10-06", "1", "Carlos Mendoza", "cmendoza@gmail.com", "987654321", "945112233",
     "Av. Javier Prado 1240, San Isidro - Lima", "Tarjeta de Debito", "PENDIENTE", "1003",
     "Camiseta Nike Dri-FIT Academy", "M", "Azul Marino", 2, 129.90, "30",
     "Indumentaria Deportiva", "Polos, casacas y pantalones Dri-FIT", "3", "2026-10-06",
     "PROGRAMADO", "1003", "Camiseta Nike Dri-FIT Academy", "M", "Azul Marino", 40, "200",
     "CD Trujillo Norte", "Parque Industrial La Esperanza - Trujillo", 200, 30, "503",
     "Nike Store Mallplaza Trujillo", "Av. America Oeste 750 - Trujillo"],
]


def hoja_1fn(wb):
    ws = wb.create_sheet("1FN")
    n = len(COLS_1FN)
    f = banner(ws, 1, "TABLAS 1FN", n)
    f = banda(ws, f + 1, "TABLA GENERAL EN 1FN", n)
    claves = [i in (0, 2, 5, 6, 10, 12, 13, 16, 19, 22, 24, 25, 27, 32) for i in range(n)]
    anchos = [11, 11, 10, 16, 22, 12, 14, 32, 17, 13, 11, 26, 11, 12, 12, 16, 11, 20, 26, 12,
              14, 14, 20, 26, 12, 13, 17, 11, 16, 32, 11, 11, 10, 22, 32]
    f = tabla(ws, f, COLS_1FN, FILAS_1FN, claves, anchos=anchos,
              rojo_cols=(5, 6), salmon_cols=(11, 12))
    ws.freeze_panes = "E4"
    return ws


# ------------------------------------------------------------------ hoja 2FN
def hoja_2fn(wb):
    ws = wb.create_sheet("2FN")
    f = banner(ws, 1, "TABLAS 2FN", 12)
    f += 1

    f = banda(ws, f, "PEDIDO", 12)
    f = tabla(ws, f,
              ["id_pedido (PK)", "fecha", "metodo_pago", "estado_envio", "CLIENTE_id_cliente (FK)",
               "SUMINISTRO_id_suministro (FK)", "nombre_cliente", "correo_electronico", "telefono",
               "telefono_secundario", "direccion"],
              [["9001", "2026-10-05", "Tarjeta de Credito", "ENTREGADO", "1", "1", "Carlos Mendoza",
                "cmendoza@gmail.com", "987654321", "945112233",
                "Av. Javier Prado 1240, San Isidro - Lima"],
               ["9002", "2026-10-06", "Yape", "ENVIADO", "2", "2", "Lucia Paredes", "lparedes@hotmail.com",
                "912345678", "NULL", "Av. Ejercito 710, Cayma - Arequipa"],
               ["9003", "2026-10-06", "Tarjeta de Debito", "PENDIENTE", "1", "3", "Carlos Mendoza",
                "cmendoza@gmail.com", "987654321", "945112233",
                "Av. Javier Prado 1240, San Isidro - Lima"]],
              [True, False, False, False, True, True] + [False] * 5,
              anchos=[12, 12, 18, 13, 18, 24, 17, 24, 13, 15, 34])

    f = banda(ws, f, "DETALLE_VENTA", 12)
    f = tabla(ws, f,
              ["PEDIDO_id_pedido (PK)", "PRODUCTO_id_producto (PK)", "talla_elegida (PK)",
               "color_elegido (PK)", "cantidad", "precio_unitario_aplicado"],
              [["9001", "1001", "42", "Negro", 1, 549.90],
               ["9001", "1002", "41", "Blanco", 1, 499.90],
               ["9002", "1002", "41", "Blanco", 1, 499.90],
               ["9003", "1003", "M", "Azul Marino", 2, 129.90]],
              [True, True, True, True, False, False],
              anchos=[17, 19, 13, 15, 10, 20])

    f = banda(ws, f, "PRODUCTO", 12)
    f = tabla(ws, f,
              ["id_producto (PK)", "nombre", "especificaciones", "precio_unitario_base", "estado",
               "CATEGORIA_id_categoria (FK)", "nombre_categoria", "descripcion_categoria"],
              [["1001", "Nike Air Zoom Pegasus 40", "Amortiguacion React con doble Zoom Air", 549.90,
                "DISPONIBLE", "10", "Calzado Running", "Zapatillas de entrenamiento y maraton"],
               ["1002", "Nike Air Force 1 07", "Capellada de cuero con suela Air", 499.90,
                "DISPONIBLE", "20", "Calzado Urbano", "Zapatillas de uso diario Sportswear"],
               ["1003", "Camiseta Nike Dri-FIT Academy", "Tejido transpirable 100% poliester", 129.90,
                "DISPONIBLE", "30", "Indumentaria Deportiva",
                "Polos, casacas y pantalones Dri-FIT"]],
              [True, False, False, False, False, True, False, False],
              anchos=[13, 28, 34, 18, 13, 20, 22, 34])

    f = banda(ws, f, "SUMINISTRO", 12)
    f = tabla(ws, f,
              ["id_suministro (PK)", "fecha_suministro", "estado_despacho",
               "ALMACEN_id_almacen (FK)", "LOCAL_id_local (FK)"],
              [["1", "2026-10-04", "ENTREGADO", "100", "501"],
               ["2", "2026-10-05", "EN RUTA", "100", "502"],
               ["3", "2026-10-06", "PROGRAMADO", "200", "503"]],
              [True, False, False, True, True],
              anchos=[16, 16, 17, 19, 15])

    f = banda(ws, f, "DETALLE_SUMINISTRO", 12)
    f = tabla(ws, f,
              ["SUMINISTRO_id_suministro (PK)", "PRODUCTO_id_producto (PK)", "talla (PK)",
               "color (PK)", "cantidad"],
              [["1", "1001", "42", "Negro", 20],
               ["1", "1002", "41", "Blanco", 15],
               ["2", "1002", "41", "Blanco", 25],
               ["3", "1003", "M", "Azul Marino", 40]],
              [True, True, True, True, False],
              anchos=[25, 22, 12, 15, 12])

    f = banda(ws, f, "ALMACEN", 12)
    f = tabla(ws, f, ["id_almacen (PK)", "nombre_almacen", "ubicacion"],
              [["100", "CD Lima Norte", "Av. Ramiro Priale Km 8 - Lima"],
               ["200", "CD Trujillo Norte", "Parque Industrial La Esperanza - Trujillo"]],
              [True, False, False], anchos=[14, 20, 36])

    f = banda(ws, f, "LOCAL", 12)
    f = tabla(ws, f, ["id_local (PK)", "nombre_local", "direccion_local"],
              [["501", "Nike Store Jockey Plaza", "Av. Javier Prado Este 4200, Surco - Lima"],
               ["502", "Nike Store Mall Arequipa", "Av. Porongoche 500, Paucarpata - Arequipa"],
               ["503", "Nike Store Mallplaza Trujillo", "Av. America Oeste 750 - Trujillo"]],
              [True, False, False], anchos=[12, 28, 38])

    f = banda(ws, f, "STOCK_ALMACEN", 12)
    f = tabla(ws, f,
              ["ALMACEN_id_almacen (PK/FK)", "PRODUCTO_id_producto (PK/FK)", "talla (PK)",
               "color (PK)", "stock_actual", "stock_minimo"],
              [["100", "1001", "42", "Negro", 85, 15],
               ["100", "1002", "41", "Blanco", 120, 20],
               ["200", "1003", "M", "Azul Marino", 200, 30]],
              [True, True, True, True, False, False],
              anchos=[22, 24, 12, 15, 13, 13])
    return ws


# ------------------------------------------------------------------ hoja 3FN
def hoja_3fn(wb):
    ws = wb.create_sheet("3FN")
    f = banner(ws, 1, "TABLAS 3FN", 8)
    f += 1
    for tab in M.ORDER:
        d = M.TABLES[tab]
        f = banda(ws, f, tab, 8)
        f = tabla(ws, f, cabecera_3fn(tab), d["filas"], columnas_clave(tab))
    for j, w in enumerate([24, 20, 18, 16, 20, 22, 14, 12], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    return ws


def construir(destino):
    wb = Workbook()
    wb.remove(wb.active)
    hoja_0fn(wb)
    hoja_1fn(wb)
    hoja_2fn(wb)
    hoja_3fn(wb)
    for ws in wb:
        ws.sheet_view.showGridLines = False
    wb.properties.creator = "Luis Coloma"
    wb.properties.lastModifiedBy = "Luis Coloma"
    wb.properties.title = "Normalizacion BD NIKE Inc."
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    wb.save(destino)
    return destino


if __name__ == "__main__":
    print(construir(M.PROJ + "/normalizacion/NORMALIZACION_BD_NIKE_GRUPO_5.xlsx"))
