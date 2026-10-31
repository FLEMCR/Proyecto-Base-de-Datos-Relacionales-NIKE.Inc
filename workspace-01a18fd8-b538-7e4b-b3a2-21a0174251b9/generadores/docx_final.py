# -*- coding: utf-8 -*-
"""Diccionario de datos (.docx) del modelo final: 10 tablas, 12 FK."""
import docx
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import modelo_final as M

AZUL = RGBColor(0x1F, 0x4E, 0x79)
ROJO = RGBColor(0xC0, 0x00, 0x00)
GRIS = RGBColor(0x5E, 0x67, 0x6E)

MOTIVO = {
    "telefono_secundario": "No todos los clientes dejan un segundo numero de contacto.",
    "descripcion": "Una categoria puede crearse solo con su nombre.",
    "especificaciones": "No todos los modelos tienen la ficha tecnica cargada.",
}


def sombrear(celda, hexcolor):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def construir(path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.9)
    sec.top_margin = Cm(1.9)
    sec.bottom_margin = Cm(1.7)

    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(9.5)
    st.paragraph_format.space_after = Pt(4)

    def h1(t):
        p = doc.add_paragraph()
        r = p.add_run(t)
        r.bold = True
        r.font.size = Pt(13)
        r.font.color.rgb = RGBColor(0x0E, 0x0E, 0x0E)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)

    def h2(t):
        p = doc.add_paragraph()
        r = p.add_run(t)
        r.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = AZUL
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)

    def par(t, size=9.5, bold=False, color=None, space=2):
        p = doc.add_paragraph()
        r = p.add_run(t)
        r.bold = bold
        r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
        p.paragraph_format.space_after = Pt(space)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return p

    def encabezado(t, cols):
        for i, e in enumerate(cols):
            c = t.rows[0].cells[i]
            c.text = ""
            r = c.paragraphs[0].add_run(e)
            r.bold = True
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            sombrear(c, "1F4E79")

    # ------------------------------------------------------------- portada
    p = doc.add_paragraph()
    r = p.add_run("DICCIONARIO DE DATOS")
    r.bold = True
    r.font.size = Pt(24)
    p.paragraph_format.space_after = Pt(0)

    p = doc.add_paragraph()
    r = p.add_run("Proyecto parcial - Base de datos NIKE Inc. (Peru)")
    r.font.size = Pt(13)
    r.bold = True
    r.font.color.rgb = RGBColor(0x44, 0x4C, 0x53)
    p.paragraph_format.space_after = Pt(14)

    datos = [("Curso", M.PROYECTO["curso"]), ("Institucion", M.PROYECTO["institucion"]),
             ("Docente", M.PROYECTO["docente"]), ("Grupo", M.PROYECTO["grupo"]),
             ("Integrantes", M.PROYECTO["integrantes"]),
             ("Motor", "Microsoft SQL Server (BD_NIKE_INC_SQL_SERVER.sql) y Oracle (BD_NIKE_INC_ORACLE.sql)")]
    t = doc.add_table(rows=len(datos), cols=2)
    t.style = "Table Grid"
    for i, (k, v) in enumerate(datos):
        c0, c1 = t.rows[i].cells
        c0.text = ""
        c1.text = ""
        r0 = c0.paragraphs[0].add_run(k)
        r0.bold = True
        r0.font.size = Pt(9.5)
        r1 = c1.paragraphs[0].add_run(v)
        r1.font.size = Pt(9.5)
        sombrear(c0, "F2F2F2")
        c0.width = Cm(4.6)
        c1.width = Cm(12.4)

    par("")
    par("Como leer este documento", 10.5, bold=True)
    par("Describe las 10 tablas del modelo relacional obtenidas en el Excel de normalizacion (0FN, 1FN, 2FN y 3FN). "
        "Para cada columna se indica el tipo, la longitud, la clave a la que pertenece, si es obligatoria y la regla "
        "de negocio que la respalda. Ninguna tabla guarda valores calculados: el total de la linea y el IGV del pedido "
        "se obtienen de la vista VW_PEDIDO_TOTALES.")

    par("Convenciones", 10.5, bold=True)
    par("PK   clave primaria de la tabla.", 9.5)
    par("FK   clave foranea: apunta a la clave primaria de otra tabla.", 9.5)
    par("PFK  columna que forma parte de la clave primaria compuesta y ademas es clave foranea.", 9.5)
    par("Obligatorio SI = NOT NULL; NO = la columna acepta NULL.", 9.5)

    # ---------------------------------------------------- 1. resumen tablas
    h1("1. Resumen de las 10 tablas")
    t = doc.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    encabezado(t, ["Nro", "Tabla", "Clasificacion", "Clave primaria", "Proposito"])
    for i, tab in enumerate(M.ORDER, 1):
        d = M.TABLES[tab]
        fila = t.add_row().cells
        vals = [str(i), tab, d["rol"], ", ".join(M.pk_cols(tab)), d["proposito"]]
        for j, v in enumerate(vals):
            fila[j].text = ""
            r = fila[j].paragraphs[0].add_run(v)
            r.font.size = Pt(8.5)
            if j == 1:
                r.bold = True
    for row in t.rows:
        for i, w in enumerate([1.0, 3.0, 2.4, 4.4, 6.4]):
            row.cells[i].width = Cm(w)

    # ------------------------------------------------- detalle por tabla
    h1("2. Detalle de columnas por tabla")
    for n, tab in enumerate(M.ORDER, 1):
        d = M.TABLES[tab]
        h2("2.%d  %s" % (n, tab))
        par(d["proposito"])
        t = doc.add_table(rows=1, cols=6)
        t.style = "Table Grid"
        encabezado(t, ["Columna", "Tipo", "Longitud", "Clave", "Obligatorio", "Referencia"])
        for col in d["cols"]:
            nombre, tipo, largo, marca, opc = col
            tx, tl = M.tipo_texto(col)
            ref = ""
            if marca in ("FK", "PFK"):
                padre, pkc = M.fk_de(tab, nombre)
                ref = "%s.%s" % (padre, pkc)
            elif marca in ("PK",):
                ref = "clave primaria"
            fila = t.add_row().cells
            vals = [nombre, tx, tl, marca if marca else "-", "NO" if opc else "SI", ref]
            for j, v in enumerate(vals):
                fila[j].text = ""
                r = fila[j].paragraphs[0].add_run(v)
                r.font.size = Pt(8)
                if j == 0:
                    r.bold = True
                if j == 3 and marca in ("PK", "PFK"):
                    r.font.color.rgb = ROJO
                    r.bold = True
                if j == 3 and marca == "FK":
                    r.font.color.rgb = RGBColor(0xA8, 0x54, 0x0B)
                    r.bold = True
                if j == 4 and opc:
                    r.font.color.rgb = RGBColor(0xA8, 0x54, 0x0B)
        for row in t.rows:
            for i, w in enumerate([4.2, 1.8, 1.8, 1.3, 2.0, 5.3]):
                row.cells[i].width = Cm(w)

    # ------------------------------------------------------ 3. longitudes
    h1("3. Por que cada campo tiene esa longitud")
    par("La longitud se calculo con el dato real que se va a guardar, no con el valor por defecto de la herramienta.")
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    encabezado(t, ["Campo", "Tipo definitivo", "Justificacion"])
    for campo, tipo, just in M.JUSTIF:
        fila = t.add_row().cells
        for j, v in enumerate([campo, tipo, just]):
            fila[j].text = ""
            r = fila[j].paragraphs[0].add_run(v)
            r.font.size = Pt(8)
            if j in (0, 1):
                r.bold = True
    for row in t.rows:
        for i, w in enumerate([5.6, 3.4, 8.2]):
            row.cells[i].width = Cm(w)

    # ------------------------------------------------- 4. columnas NULL
    h1("4. Columnas que admiten NULL")
    opc = M.opcionales()
    par("Solo %d columnas del modelo aceptan un valor vacio. Todo lo demas es NOT NULL." % len(opc))
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    encabezado(t, ["Tabla", "Columna", "Por que puede ir vacia"])
    for tabla, col in opc:
        fila = t.add_row().cells
        for j, v in enumerate([tabla, col, MOTIVO.get(col, "")]):
            fila[j].text = ""
            r = fila[j].paragraphs[0].add_run(v)
            r.font.size = Pt(8)
            if j in (0, 1):
                r.bold = True
    for row in t.rows:
        for i, w in enumerate([4.4, 4.6, 8.2]):
            row.cells[i].width = Cm(w)

    # ------------------------------------------------------- 5. catalogo FK
    h1("5. Catalogo de claves foraneas (12 restricciones)")
    t = doc.add_table(rows=1, cols=6)
    t.style = "Table Grid"
    encabezado(t, ["Restriccion", "Tabla hija (N)", "Columna FK", "Tabla padre (1)", "Columna PK", "Cardinalidad"])
    for nombre, hija, col, padre, pkc in M.FKS:
        fila = t.add_row().cells
        for j, v in enumerate([nombre, hija, col, padre, pkc, "1:N"]):
            fila[j].text = ""
            r = fila[j].paragraphs[0].add_run(v)
            r.font.size = Pt(8)
            if j == 0:
                r.bold = True
    for row in t.rows:
        for i, w in enumerate([4.4, 3.1, 3.4, 2.9, 2.2, 1.4]):
            row.cells[i].width = Cm(w)
    par("")
    par("Las 8 relaciones del MER quedan resueltas asi: las de uno a muchos (REALIZA, CLASIFICA, RECIBE, ABASTECE y "
        "PROVEE) se convierten en claves foraneas, y las de muchos a muchos (CONTIENE, ALMACENA e INCLUYE) en las "
        "tablas DETALLE_VENTA, STOCK_ALMACEN y DETALLE_SUMINISTRO.", 9.0)

    # ------------------------------------------- 6. calculados y vista
    h1("6. Datos calculados: no se guardan")
    par("El total de la linea y el IGV del pedido se pueden obtener de otros datos, asi que no se almacenan.")
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    encabezado(t, ["Elemento", "Formula", "Donde vive"])
    for fila_txt in [
        ("total de la linea", "cantidad x precio_unitario_aplicado", "se calcula en la vista VW_PEDIDO_TOTALES"),
        ("VW_PEDIDO_TOTALES.total_pedido", "suma de los totales del pedido", "vista: no se almacena"),
        ("VW_PEDIDO_TOTALES.igv", "total_pedido x 0.18", "vista: no se almacena"),
        ("VW_PEDIDO_TOTALES.total_con_igv", "total_pedido x 1.18", "vista: no se almacena"),
    ]:
        fila = t.add_row().cells
        for j, v in enumerate(fila_txt):
            fila[j].text = ""
            r = fila[j].paragraphs[0].add_run(v)
            r.font.size = Pt(8)
            if j == 0:
                r.bold = True
    for row in t.rows:
        for i, w in enumerate([5.0, 5.0, 7.2]):
            row.cells[i].width = Cm(w)

    # ---------------------------------------------------- 7. estados
    h1("7. Catalogo de estados")
    par("El pedido y el suministro son dos procesos distintos, por eso cada uno tiene su propio estado.")
    t = doc.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    encabezado(t, ["Campo", "Tabla", "Para que sirve", "Valores permitidos", "Ejemplo"])
    for fila_txt in M.ESTADOS:
        fila = t.add_row().cells
        for j, v in enumerate(fila_txt):
            fila[j].text = ""
            r = fila[j].paragraphs[0].add_run(v)
            r.font.size = Pt(8)
            if j == 0:
                r.bold = True
    for row in t.rows:
        for i, w in enumerate([3.0, 2.4, 4.0, 3.6, 4.2]):
            row.cells[i].width = Cm(w)

    doc.core_properties.author = "Luis Coloma"
    doc.core_properties.title = "Diccionario de datos - Proyecto BD NIKE Inc. (Grupo 05)"
    doc.core_properties.comments = ""
    doc.core_properties.last_modified_by = "Luis Coloma"
    doc.save(path)
    return path


if __name__ == "__main__":
    print(construir(M.PROJ + "/DICCIONARIO_DE_DATOS_NIKE_GRUPO_5.docx"))
