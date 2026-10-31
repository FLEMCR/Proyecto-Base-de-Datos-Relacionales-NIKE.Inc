# -*- coding: utf-8 -*-
"""Renderizador simple de la presentacion a PNG/PDF (sin LibreOffice).

Soporta lo que usa el deck: imagenes, rectangulos y rectangulos redondeados,
lineas, cuadros de texto con vinetas y tablas.
"""
import io
import os

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

F = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FB = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"

EMU_W, EMU_H = 12192000, 6858000
PX_W, PX_H = 1920, 1080
ESC = PX_W / EMU_W


def fuente(tam_pt, negrita):
    # 0.93 compensa que DejaVu es mas ancha que Arial, que es la fuente del deck
    px = max(6, int(round(tam_pt * PX_W / 960.0 * 0.93)))
    return ImageFont.truetype(FB if negrita else F, px)


def emu_px(v):
    return v * ESC


def color_de(fill):
    try:
        if fill.type is not None and fill.fore_color.rgb is not None:
            return tuple(int(str(fill.fore_color.rgb)[i:i + 2], 16) for i in (0, 2, 4))
    except Exception:
        pass
    return None


def color_linea(sh, default=(0, 0, 0)):
    try:
        if sh.line.color and sh.line.color.rgb is not None:
            return tuple(int(str(sh.line.color.rgb)[i:i + 2], 16) for i in (0, 2, 4))
    except Exception:
        pass
    return default


def ancho_linea(sh):
    try:
        if sh.line.width:
            return max(1, int(emu_px(sh.line.width)))
    except Exception:
        pass
    return 1


A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def texto_de_tf(tf):
    """Devuelve [(texto, tam, negrita, color, alineacion, vinneta)] por parrafo."""
    salida = []
    for p in tf.paragraphs:
        # los saltos de linea (a:br) se conservan como \n
        trozos = []
        for el in p._p:
            if el.tag == A + "r":
                t = el.find(A + "t")
                trozos.append(t.text if t is not None and t.text else "")
            elif el.tag == A + "br":
                trozos.append("\n")
        txt = "".join(trozos)
        if not txt.strip():
            txt = "".join(r.text for r in p.runs)
        if not txt.strip():
            salida.append(None)
            continue
        run = p.runs[0]
        tam = run.font.size.pt if run.font.size else 12.0
        negrita = bool(run.font.bold)
        try:
            col = run.font.color.rgb
            col = tuple(int(str(col)[i:i + 2], 16) for i in (0, 2, 4)) if col else (20, 20, 20)
        except Exception:
            col = (20, 20, 20)
        alin = str(p.alignment) if p.alignment else "LEFT"
        pPr = p._p.find("{http://schemas.openxmlformats.org/drawingml/2006/main}pPr")
        vinneta = None
        if pPr is not None:
            bu = pPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar")
            if bu is not None:
                vinneta = bu.get("char")
        salida.append((txt, tam, negrita, col, alin, vinneta))
    return salida


def envolver(d, txt, f, ancho):
    palabras = txt.split(" ")
    lineas, actual = [], ""
    for w in palabras:
        prueba = (actual + " " + w).strip()
        if d.textlength(prueba, font=f) <= ancho or not actual:
            actual = prueba
        else:
            lineas.append(actual)
            actual = w
    if actual:
        lineas.append(actual)
    return lineas


def dibujar_texto(d, tf, x, y, w, h, margen=3):
    parrafos = texto_de_tf(tf)
    yy = y + margen
    for item in parrafos:
        if item is None:
            yy += fuente(12, False).size * 0.9
            continue
        txt, tam, negrita, col, alin, vinneta = item
        f = fuente(tam, negrita)
        sangria = 0
        if vinneta:
            sangria = f.size * 1.1
        lineas = []
        for trozo in txt.split("\n"):
            if trozo == "":
                lineas.append("")
            else:
                lineas.extend(envolver(d, trozo, f, max(10, w - 2 * margen - sangria)))
        for i, linea in enumerate(lineas):
            if "CENTER" in alin.upper():
                tw = d.textlength(linea, font=f)
                tx = x + (w - tw) / 2
            elif "RIGHT" in alin.upper():
                tw = d.textlength(linea, font=f)
                tx = x + w - margen - tw
            else:
                tx = x + margen + (sangria if i == 0 or vinneta else sangria)
            if i == 0 and vinneta:
                # cuadradito verde en el margen de la vineta
                lado = max(5, int(f.size * 0.40))
                bx = x + margen
                d.rectangle([bx, yy + f.size * 0.34, bx + lado, yy + f.size * 0.34 + lado],
                            fill=(0x8B, 0xC3, 0x4A))
                tx = x + margen + sangria
            d.text((tx, yy), linea, font=f, fill=col)
            yy += int(f.size * 1.12)
    return yy


def render_slide(slide, ruta):
    img = Image.new("RGB", (PX_W, PX_H), "white")
    d = ImageDraw.Draw(img)
    for sh in slide.shapes:
        x, y = emu_px(sh.left or 0), emu_px(sh.top or 0)
        w, h = emu_px(sh.width or 0), emu_px(sh.height or 0)
        tipo = sh.shape_type
        if tipo == MSO_SHAPE_TYPE.PICTURE:
            try:
                im = Image.open(io.BytesIO(sh.image.blob)).convert("RGB")
                im = im.resize((max(1, int(w)), max(1, int(h))), Image.LANCZOS)
                img.paste(im, (int(x), int(y)))
            except Exception:
                pass
            continue
        if tipo == MSO_SHAPE_TYPE.LINE:
            flip = sh._element.spPr.find(
                "{http://schemas.openxmlformats.org/drawingml/2006/main}xfrm")
            fh = flip is not None and flip.get("flipH") == "1"
            fv = flip is not None and flip.get("flipV") == "1"
            x1, y1 = (x + w, y) if fh else (x, y)
            x2, y2 = (x, y) if fh else (x + w, y + h)
            if fv:
                y1, y2 = y2, y1
            d.line([x1, y1, x2, y2], fill=color_linea(sh, (0xDD, 0xDD, 0xDD)), width=ancho_linea(sh))
            continue
        if getattr(sh, "has_table", False) and sh.has_table:
            tabla = sh.table
            anchos = [emu_px(c.width) for c in tabla.columns]
            yy = y
            for i, fila in enumerate(tabla.rows):
                alto = emu_px(fila.height)
                xx = x
                for j, celda in enumerate(fila.cells):
                    relleno = color_de(celda.fill)
                    if relleno:
                        d.rectangle([xx, yy, xx + anchos[j], yy + alto], fill=relleno)
                    if celda.text_frame.text.strip():
                        dibujar_texto(d, celda.text_frame, xx, yy, anchos[j], alto, margen=4)
                    xx += anchos[j]
                yy += alto
            continue
        # autoshapes y cuadros de texto
        geometria = ""
        try:
            geometria = str(sh.auto_shape_type)
        except Exception:
            geometria = ""
        relleno = color_de(sh.fill) if hasattr(sh, "fill") else None
        if "OVAL" in geometria:
            if relleno:
                d.ellipse([x, y, x + w, y + h], fill=relleno, outline=color_linea(sh),
                          width=ancho_linea(sh))
        elif "ROUNDED" in geometria:
            radio = min(w, h) * 0.25
            if relleno:
                d.rounded_rectangle([x, y, x + w, y + h], radius=radio, fill=relleno,
                                    outline=color_linea(sh, (0xE2, 0xE5, 0xE8)), width=ancho_linea(sh))
            else:
                d.rounded_rectangle([x, y, x + w, y + h], radius=radio,
                                    outline=color_linea(sh, (0xE2, 0xE5, 0xE8)), width=ancho_linea(sh))
        elif "RECTANGLE" in geometria and relleno:
            d.rectangle([x, y, x + w, y + h], fill=relleno, outline=color_linea(sh),
                        width=ancho_linea(sh))
        if sh.has_text_frame and sh.text_frame.text.strip():
            dibujar_texto(d, sh.text_frame, x, y, w, h)
    img.save(ruta)
    return ruta


def main(pptx, carpeta):
    os.makedirs(carpeta, exist_ok=True)
    prs = Presentation(pptx)
    rutas = []
    for i, s in enumerate(prs.slides, 1):
        r = os.path.join(carpeta, "lamina_%02d.png" % i)
        render_slide(s, r)
        rutas.append(r)
    return rutas


def a_pdf(rutas, salida):
    import pymupdf
    doc = pymupdf.open()
    for r in rutas:
        im = Image.open(r)
        pg = doc.new_page(width=im.width * 0.5, height=im.height * 0.5)
        pg.insert_image(pg.rect, filename=r)
    doc.save(salida)
    return salida


if __name__ == "__main__":
    base = "/home/user/proyecto_parcial/PRESENTACION_PROYECTO_BD_NIKE_GRUPO_5"
    rutas = main(base + ".pptx", "/tmp/laminas")
    print(len(rutas), "laminas")
    print(a_pdf(rutas, "/tmp/PRESENTACION_nueva.pdf"))
