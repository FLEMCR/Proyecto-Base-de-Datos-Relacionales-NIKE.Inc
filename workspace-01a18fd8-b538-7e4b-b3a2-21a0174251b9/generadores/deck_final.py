# -*- coding: utf-8 -*-
"""Actualiza la presentacion: laminas 4 y 5 rehechas y diagramas nuevos en 9, 10 y 11.

Las laminas 4 y 5 estaban rasterizadas con el MER anterior (DESPACHA, Total calculado).
Aqui se rehacen como laminas normales de PowerPoint (texto Arial, igual que el resto).
"""
import copy

from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from lxml import etree
from PIL import Image

W, H = 12192000, 6858000
VERDE = RGBColor(0x8B, 0xC3, 0x4A)
NEGRO = RGBColor(0x11, 0x11, 0x14)
GRIS = RGBColor(0x5C, 0x63, 0x68)
GRIS_CLARO = RGBColor(0x6E, 0x73, 0x78)


def frac(x, y, w, h):
    return Emu(int(x * W)), Emu(int(y * H)), Emu(int(w * W)), Emu(int(h * H))


def caja_texto(slide, x, y, w, h, texto, tam=12, negrita=False, color=NEGRO, align=PP_ALIGN.LEFT,
               espacio=0):
    sh = slide.shapes.add_textbox(*frac(x, y, w, h))
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = texto
    f = r.font
    f.name = "Arial"
    f.size = Pt(tam)
    f.bold = negrita
    f.color.rgb = color
    p.space_after = Pt(espacio)
    return sh


def vineta(slide, x, y, w, h, texto, tam=11.5, color=NEGRO):
    sh = slide.shapes.add_textbox(*frac(x, y, w, h))
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Emu(Emu(int(0.028 * W))).emu if False else 0)) if False else "228600")
    pPr.set("indent", "-228600")
    buClr = etree.SubElement(pPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}buClr")
    etree.SubElement(buClr, "{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr").set("val", "8BC34A")
    buFont = etree.SubElement(pPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}buFont")
    buFont.set("typeface", "Arial")
    buChar = etree.SubElement(pPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}buChar")
    buChar.set("char", "\u25a0")
    r = p.add_run()
    r.text = texto
    r.font.name = "Arial"
    r.font.size = Pt(tam)
    r.font.color.rgb = color
    return sh


def pastilla(slide, x, y, w, h, texto):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, *frac(x, y, w, h))
    sh.adjustments[0] = 0.5
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor(0xF4, 0xF5, 0xF6)
    sh.line.fill.background()
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = texto
    r.font.name = "Arial"
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.color.rgb = GRIS_CLARO
    return sh


def tarjeta(slide, x, y, w, h):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, *frac(x, y, w, h))
    sh.adjustments[0] = 0.02
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    sh.line.color.rgb = RGBColor(0xE2, 0xE5, 0xE8)
    sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    return sh


def imagen_en(slide, ruta, x, y, w, h, margen=0.006):
    """Inserta la imagen centrada dentro del rectangulo, conservando la proporcion."""
    im = Image.open(ruta)
    ar = im.width / im.height
    cw, ch = w * (1 - 2 * margen), h * (1 - 2 * margen)
    if cw / ch > ar:
        nh = ch
        nw = ch * ar
    else:
        nw = cw
        nh = cw / ar
    nx = x + (w - nw) / 2
    ny = y + (h - nh) / 2
    return slide.shapes.add_picture(ruta, *frac(nx, ny, nw, nh))


def limpiar(slide):
    for sh in list(slide.shapes):
        sh._element.getparent().remove(sh._element)


def logo(slide, blob):
    """Reusa el logo IDAT de la esquina superior derecha."""
    from io import BytesIO
    slide.shapes.add_picture(BytesIO(blob), *frac(0.8917, 0.0444, 0.0396, 0.1162))


def linea_h(slide, x1, x2, y, color=RGBColor(0xDD, 0xDD, 0xDD)):
    from pptx.enum.shapes import MSO_CONNECTOR
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(int(x1 * W)), Emu(int(y * H)),
                                    Emu(int(x2 * W)), Emu(int(y * H)))
    ln.line.color.rgb = color
    ln.line.width = Pt(0.9)
    return ln


def pie(slide, numero, total=16):
    linea_h(slide, 0.06, 0.95, 0.9375)
    caja_texto(slide, 0.06, 0.948, 0.60, 0.03,
               "Proyecto parcial  |  Disenio de Base de Datos Relacionales  |  Grupo 05",
               tam=9, color=GRIS)
    caja_texto(slide, 0.90, 0.948, 0.04, 0.03, "%02d / %d" % (numero, total), tam=9, color=GRIS,
               align=PP_ALIGN.RIGHT)


def reconstruir_slide4(slide, logo_blob, mer_png):
    limpiar(slide)
    logo(slide, logo_blob)
    tarjeta(slide, 0.049, 0.111, 0.408, 0.625)
    imagen_en(slide, mer_png, 0.052, 0.118, 0.402, 0.611, margen=0.004)
    pastilla(slide, 0.490, 0.088, 0.078, 0.030, "CAPITULO 3")
    caja_texto(slide, 0.490, 0.118, 0.385, 0.105, "Modelo Entidad-Relacion (MER)", tam=26,
               negrita=True)
    caja_texto(slide, 0.490, 0.240, 0.385, 0.070,
               "Notacion Chen en Draw.io: no depende de ningun motor, solo muestra las entidades "
               "del negocio y como se relacionan.", tam=11.5, color=GRIS)
    vinetas = [
        "7 entidades: CLIENTE, PEDIDO, PRODUCTO, CATEGORIA, ALMACEN, LOCAL y SUMINISTRO.",
        "8 relaciones: REALIZA, CONTIENE, CLASIFICA, ALMACENA, PROVEE, ABASTECE, RECIBE e INCLUYE.",
        "Las 5 relaciones 1:N llevan la flecha hacia el lado 1; los rombos M:N van sin flecha.",
        "CONTIENE guarda talla, color, cantidad, precio aplicado y el almacen de origen.",
        "Ningun atributo calculado: no hay Total ni IGV en el diagrama.",
        "Telefono_secundario va punteado: es el unico atributo opcional del cliente.",
    ]
    y = 0.330
    for v in vinetas:
        vineta(slide, 0.490, y, 0.390, 0.075, v)
        y += 0.095
    pie(slide, 4)


def reconstruir_slide5(slide, logo_blob, mapa_png):
    limpiar(slide)
    logo(slide, logo_blob)
    pastilla(slide, 0.049, 0.075, 0.078, 0.030, "CAPITULO 3")
    caja_texto(slide, 0.049, 0.105, 0.700, 0.060, "Las 8 relaciones del MER y como se leen",
               tam=26, negrita=True)
    caja_texto(slide, 0.049, 0.180, 0.900, 0.030,
               "Cada rombo se lee como una frase en espanol. El numero del rombo coincide con la "
               "lista de la derecha.", tam=11.5, color=GRIS)
    tarjeta(slide, 0.049, 0.225, 0.600, 0.405)
    imagen_en(slide, mapa_png, 0.052, 0.230, 0.594, 0.395, margen=0.004)
    caja_texto(slide, 0.049, 0.645, 0.600, 0.070,
               "Los rombos M:N (2, 4 y 7) no llevan flecha; en las 1:N la punta toca la entidad del "
               "lado 1. PROVEE e INCLUYE son las dos lineas que unen la zona de ventas con la de "
               "abastecimiento.", tam=10, color=GRIS)

    caja_texto(slide, 0.665, 0.225, 0.305, 0.030, "Lectura de cada relacion", tam=13.5,
               negrita=True)
    lectura = [
        ("REALIZA (1,N)", "un CLIENTE realiza muchos PEDIDOS"),
        ("CONTIENE (M,N)", "un PEDIDO contiene muchos PRODUCTOS (talla, color, cantidad, almacen)"),
        ("CLASIFICA (1,N)", "una CATEGORIA clasifica muchos PRODUCTOS"),
        ("ALMACENA (M,N)", "un ALMACEN almacena muchos PRODUCTOS por talla y color"),
        ("PROVEE (1,N)", "un SUMINISTRO provee a muchos PEDIDOS"),
        ("ABASTECE (1,N)", "un ALMACEN abastece muchos SUMINISTROS"),
        ("INCLUYE (M,N)", "un SUMINISTRO incluye muchos PRODUCTOS (talla, color, cantidad)"),
        ("RECIBE (1,N)", "un LOCAL recibe muchos SUMINISTROS"),
    ]
    y = 0.258
    for i, (titulo, frase) in enumerate(lectura, 1):
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, *frac(0.665, y + 0.004, 0.017, 0.030))
        circ.fill.solid()
        circ.fill.fore_color.rgb = RGBColor(0xD4, 0xE8, 0x4A)
        circ.line.color.rgb = RGBColor(0x8A, 0x9B, 0x1E)
        circ.line.width = Pt(0.75)
        circ.shadow.inherit = False
        tf = circ.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = str(i)
        r.font.name = "Arial"
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = NEGRO
        caja_texto(slide, 0.688, y, 0.282, 0.030, titulo, tam=11, negrita=True)
        caja_texto(slide, 0.688, y + 0.027, 0.282, 0.048, frase, tam=9.5, color=GRIS)
        linea_h(slide, 0.665, 0.970, y + 0.076, RGBColor(0xE8, 0xEA, 0xEC))
        y += 0.082
    pie(slide, 5)


def reemplazar_imagen(shape, ruta):
    """Cambia el binario de una imagen y ajusta la caja para no deformarla."""
    im = Image.open(ruta)
    ar = im.width / im.height
    cx, cy = shape.width, shape.height
    if cx / cy > ar:
        nw, nh = int(cy * ar), cy
    else:
        nw, nh = cx, int(cx / ar)
    shape.left = shape.left + (cx - nw) // 2
    shape.top = shape.top + (cy - nh) // 2
    shape.width, shape.height = nw, nh
    rid = shape._element.blipFill.blip.rEmbed
    part = shape.part.related_part(rid)
    part._blob = open(ruta, "rb").read()


def reemplazos_texto(prs):
    """Textos que quedaron con la version anterior del modelo."""
    cambios = {
        "Total queda calculado y el IGV sale de la vista":
            "Ningun dato calculado se guarda: total e IGV salen de una vista",
        "total e igv se pueden calcular":
            "total e igv no se guardan",
        "total pasa a columna calculada (cantidad x precio) e igv se obtiene con una vista; no se":
            "se eliminaron las dos columnas: el total de la linea y el igv del pedido se calculan en",
        "total calculado + VW_PEDIDO_TOTALES":
            "vista VW_PEDIDO_TOTALES",
        "Atributos: * obligatorio y  -  opcional (telefono_secundario, descripcion, especificaciones, estado_despacho y el suministro del pedido).  =  columna calculada.  10 tablas.":
            "Atributos: * obligatorio y  -  opcional (telefono_secundario, descripcion y especificaciones).  Sin columnas calculadas.  10 tablas.",
        "Solo 5 columnas admiten NULL; las claves y las foraneas nunca.":
            "Solo 3 columnas admiten NULL; las claves y las foraneas nunca.",
        "Cabecera de la compra. El IGV ya no se guarda: se calcula con la vista":
            "Cabecera de la compra. Sin columna de IGV: se calcula con la vista",
        "El total se puede calcular (cantidad x precio) y el IGV es el 18 % de la suma: guardarlos era":
            "El total se calcula (cantidad x precio) y el IGV es el 18 % de la suma: guardarlos era",
        "redundante. Total quedo como columna calculada y el IGV se obtiene con la vista":
            "redundante. Se eliminaron las dos columnas y se obtienen de la vista",
        "Antes: total e igv digitados.  Despues: calculados.":
            "Antes: total e igv digitados.  Despues: fuera de las tablas.",
        "Columna virtual (Oracle) o columna calculada PERSISTED (SQL Server) con la":
            "Ninguna: total e IGV se calculan en la vista VW_PEDIDO_TOTALES, no son columnas",
        "cantidad x precio_unitario_aplicado: no se digita, se genera sola en la base.":
            "cantidad x precio_unitario_aplicado: no se almacena, se calcula en la vista.",
    }
    tocados = 0
    for s in prs.slides:
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    for viejo, nuevo in cambios.items():
                        if viejo and viejo in r.text:
                            r.text = r.text.replace(viejo, nuevo)
                            tocados += 1
    return tocados


def quitar_filas(prs, indice_slide, textos):
    """Elimina de las tablas las filas cuya primera celda coincide con algun texto."""
    quitadas = []
    for sh in prs.slides[indice_slide].shapes:
        if not (getattr(sh, "has_table", False) and sh.has_table):
            continue
        tabla = sh.table
        for fila in list(tabla.rows):
            etiqueta = fila.cells[0].text.strip().lower()
            if any(etiqueta == t for t in textos):
                tr = fila._tr
                tr.getparent().remove(tr)
                quitadas.append(etiqueta)
    return quitadas


def fijar(prs, indice_slide, x, y, texto, tol=0.006):
    """Reemplaza el texto de la caja que esta en esa posicion (mantiene el formato)."""
    for sh in prs.slides[indice_slide].shapes:
        if not sh.has_text_frame:
            continue
        if abs(sh.left / W - x) < tol and abs(sh.top / H - y) < tol:
            tf = sh.text_frame
            p = tf.paragraphs[0]
            if not p.runs:
                continue
            p.runs[0].text = texto
            for r in p.runs[1:]:
                r.text = ""
            # se borran los parrafos sobrantes de la celda
            for extra in list(tf.paragraphs[1:]):
                extra._p.getparent().remove(extra._p)
            return True
    return False


def main():
    ruta = "/home/user/proyecto_parcial/PRESENTACION_PROYECTO_BD_NIKE_GRUPO_5.pptx"
    prs = Presentation(ruta)
    # logo de la esquina (se toma de cualquier lamina que lo tenga)
    logo_blob = None
    for sh in prs.slides[1].shapes:
        if sh.shape_type == 13 and sh.width < 600000:
            logo_blob = sh.image.blob
            break
    mer = "/home/user/proyecto_parcial/nivel_conceptual/MER_informe.png"
    mapa = "/home/user/generadores/fig_mapa_relaciones.png"
    reconstruir_slide4(prs.slides[3], logo_blob, mer)
    reconstruir_slide5(prs.slides[4], logo_blob, mapa)
    # laminas 9, 10 y 11: modelo logico, relacional y fisico
    pares = {9: "/home/user/proyecto_parcial/nivel_logico/MODELO_LOGICO.png",
             10: "/home/user/proyecto_parcial/nivel_logico/MODELO_RELACIONAL.png",
             11: "/home/user/proyecto_parcial/nivel_fisico/MODELO_FISICO.png"}
    for num, ruta_img in pares.items():
        s = prs.slides[num - 1]
        for sh in list(s.shapes):
            if sh.shape_type == 13 and sh.width > 5000000:
                reemplazar_imagen(sh, ruta_img)
    # lamina 7: la figura de la fase 3FN se rehace (mostraba total como columna calculada)
    for sh in prs.slides[6].shapes:
        if sh.shape_type == 13 and sh.width > 2000000:
            reemplazar_imagen(sh, "/home/user/generadores/fig_fase3fn.png")
            sh.left, sh.top = Emu(int(0.050 * W)), Emu(int(0.215 * H))
            sh.width = Emu(int(0.560 * W))
            sh.height = Emu(int(sh.width / (3520 / 1400)))
            break
    # lamina 8: fila 9 de la tabla de correcciones
    fijar(prs, 7, 0.327, 0.608,
          "Se eliminaron las dos columnas: el total de la linea y el igv del pedido se calculan en "
          "la vista y no se digitan.")
    # lamina 13: fila de opcionales y fila de columnas calculadas
    fijar(prs, 12, 0.165, 0.364,
          "Acepta NULL: telefono_secundario, descripcion y especificaciones.")
    fijar(prs, 12, 0.065, 0.612, "no se guarda")
    fijar(prs, 12, 0.165, 0.612, "Ningun valor calculado se guarda en las tablas.")
    fijar(prs, 12, 0.477, 0.612, "Vista VW_PEDIDO_TOTALES: calcula total, igv y total con IGV.")
    # lamina 12: el tipo de la fila del total de la linea
    fijar(prs, 11, 0.287, 0.684, "no se almacena")
    # lamina 14: PK de DETALLE_SUMINISTRO
    fijar(prs, 13, 0.263, 0.763, "suministro, producto, talla, color")
    print("filas quitadas (lamina 12):",
          quitar_filas(prs, 11, {"total (calculado)", "igv"}))
    print("textos corregidos:", reemplazos_texto(prs))
    prs.save(ruta)
    print("ok", ruta)


if __name__ == "__main__":
    main()
